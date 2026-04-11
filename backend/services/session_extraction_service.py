"""
Extração assíncrona de texto para arquivos efêmeros de sessão.

Estratégia:
- PDF: tenta PyMuPDF primeiro
- Se o PDF não tiver texto útil, cai para OCR por página
- DOCX/XLSX/TXT/CSV/MD: leitura nativa
- Imagens: OCR direto
"""

from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timezone
from pathlib import Path

from backend.services.ephemeral_rag_service import chunk_text
from backend.services.session_file_service import (
    get_file_workspace_dir,
    get_extracted_file_path,
    get_workspace_images_dir,
    get_workspace_manifest_path,
    list_session_files,
    get_stored_file_path,
    mark_file_failed,
    mark_file_processing,
    mark_file_ready,
)

logger = logging.getLogger(__name__)

PDF_NATIVE_MIN_CHARS = 80
STUCK_PROCESSING_AFTER_SECONDS = 20


def process_uploaded_file(session_id: str, file_id: str) -> None:
    try:
        entry = mark_file_processing(session_id, file_id)
        file_path = get_stored_file_path(session_id, file_id)
        extracted_text = extract_text_for_file(file_path, entry.get("mime_type", ""))
        output_path = get_extracted_file_path(session_id, file_id)
        write_extracted_text(output_path, extracted_text)
        image_paths = extract_workspace_images(file_path, session_id=session_id, file_id=file_id)
        chunk_count = len(chunk_text(extracted_text))
        write_workspace_manifest(
            session_id=session_id,
            file_id=file_id,
            original_name=str(entry.get("original_name") or file_path.name),
            mime_type=str(entry.get("mime_type") or "application/octet-stream"),
            extracted_text_path=output_path,
            image_paths=image_paths,
            text_char_count=len(extracted_text),
            chunk_count=chunk_count,
        )
        mark_file_ready(
            session_id,
            file_id,
            text_char_count=len(extracted_text),
            image_count=len(image_paths),
            chunk_count=chunk_count,
        )
        logger.info(
            "Workspace documental concluído para arquivo %s da sessão %s (%s chars, %s imagens)",
            file_id,
            session_id,
            len(extracted_text),
            len(image_paths),
        )
    except Exception as exc:
        logger.error(
            "Falha na extração do arquivo %s da sessão %s: %s",
            file_id,
            session_id,
            exc,
        )
        mark_file_failed(session_id, file_id, str(exc))


def recover_pending_session_files(session_id: str) -> list[str]:
    """
    Recupera anexos presos em `uploaded` ou `processing`.

    Isso protege o front contra spinner infinito quando a background task
    não inicia, cai no meio do caminho, ou extrai o texto mas não chega a
    atualizar o manifesto final.
    """
    recovered_file_ids: list[str] = []

    for entry in list_session_files(session_id):
        file_id = str(entry.get("file_id") or "").strip()
        status = str(entry.get("status") or "uploaded").strip().lower()
        if not file_id or status not in {"uploaded", "processing"}:
            continue

        extracted_path = get_extracted_file_path(session_id, file_id)
        if extracted_path.exists():
            mark_file_ready(session_id, file_id)
            recovered_file_ids.append(file_id)
            continue

        if status == "uploaded" or _is_stale_processing(entry):
            logger.warning(
                "Arquivo %s da sessão %s ficou preso em %s. Tentando reprocessar.",
                file_id,
                session_id,
                status,
            )
            process_uploaded_file(session_id, file_id)
            recovered_file_ids.append(file_id)

    return recovered_file_ids


def _is_stale_processing(entry: dict) -> bool:
    created_at = str(entry.get("created_at") or "").strip()
    if not created_at:
        return True

    try:
        created_dt = datetime.fromisoformat(created_at)
        if created_dt.tzinfo is None:
            created_dt = created_dt.replace(tzinfo=timezone.utc)
        created_dt = created_dt.astimezone(timezone.utc)
    except Exception:
        return True

    age_seconds = (datetime.now(timezone.utc) - created_dt).total_seconds()
    return age_seconds >= STUCK_PROCESSING_AFTER_SECONDS


def extract_text_for_file(file_path: Path, mime_type: str = "") -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        native_text = extract_pdf_text_native(file_path)
        if len(native_text.strip()) >= PDF_NATIVE_MIN_CHARS:
            return native_text
        logger.info(
            "PDF %s sem texto suficiente via PyMuPDF (%s chars). Acionando OCR.",
            file_path.name,
            len(native_text.strip()),
        )
        return extract_pdf_text_with_ocr(file_path)

    if suffix == ".docx":
        return extract_docx_text(file_path)

    if suffix == ".xlsx":
        return extract_xlsx_text(file_path)

    if suffix in {".txt", ".md", ".csv"}:
        return extract_plain_text(file_path)

    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"} or mime_type.startswith("image/"):
        return extract_image_text_with_ocr(file_path)

    return extract_plain_text(file_path)


def extract_workspace_images(file_path: Path, *, session_id: str, file_id: str) -> list[str]:
    if file_path.suffix.lower() != ".pdf":
        return []

    try:
        import fitz
    except ImportError:
        logger.warning("PyMuPDF não instalado. Extração de imagens desabilitada para %s", file_path.name)
        return []

    images_dir = get_workspace_images_dir(session_id, file_id)
    workspace_dir = get_file_workspace_dir(session_id, file_id)
    workspace_dir.mkdir(parents=True, exist_ok=True)
    images_dir.mkdir(parents=True, exist_ok=True)

    extracted: list[str] = []
    seen_xrefs: set[int] = set()
    try:
        with fitz.open(file_path) as document:
            for page_index, page in enumerate(document, start=1):
                for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                    xref = int(image_info[0])
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)
                    image_data = document.extract_image(xref)
                    image_bytes = image_data.get("image")
                    image_ext = image_data.get("ext") or "png"
                    if not image_bytes:
                        continue
                    image_path = images_dir / f"page_{page_index:03d}_image_{image_index:03d}.{image_ext}"
                    image_path.write_bytes(image_bytes)
                    extracted.append(str(image_path))
    except Exception as exc:
        logger.warning("Falha ao extrair imagens do PDF %s: %s", file_path.name, exc)
        return []

    return extracted


def write_workspace_manifest(
    *,
    session_id: str,
    file_id: str,
    original_name: str,
    mime_type: str,
    extracted_text_path: Path,
    image_paths: list[str],
    text_char_count: int,
    chunk_count: int,
) -> None:
    manifest_path = get_workspace_manifest_path(session_id, file_id)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "file_id": file_id,
        "original_name": original_name,
        "mime_type": mime_type,
        "workspace_status": "ready",
        "extracted_text_path": str(extracted_text_path),
        "image_paths": image_paths,
        "image_count": len(image_paths),
        "text_char_count": text_char_count,
        "chunk_count": chunk_count,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    manifest_path.write_text(json.dumps(payload, ensure_ascii=True, indent=2), encoding="utf-8")


def extract_pdf_text_native(file_path: Path) -> str:
    try:
        import fitz
    except ImportError:
        logger.warning("PyMuPDF não instalado. Usando fallback PyPDF2 para %s", file_path.name)
        return extract_pdf_text_pypdf2(file_path)

    texts: list[str] = []
    try:
        with fitz.open(file_path) as document:
            for page in document:
                page_text = page.get_text("text") or ""
                if page_text.strip():
                    texts.append(page_text.strip())
        return "\n\n".join(texts).strip()
    except Exception as exc:
        logger.error("Erro na leitura nativa do PDF %s: %s", file_path, exc)
        raise


def extract_pdf_text_pypdf2(file_path: Path) -> str:
    try:
        import PyPDF2
    except ImportError as exc:
        raise ValueError("PyPDF2 não instalado. Não foi possível extrair texto do PDF.") from exc

    try:
        texts: list[str] = []
        with file_path.open("rb") as file_obj:
            reader = PyPDF2.PdfReader(file_obj)
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    texts.append(page_text.strip())
        return "\n\n".join(texts).strip()
    except Exception as exc:
        logger.error("Erro no fallback PyPDF2 para %s: %s", file_path, exc)
        raise


def extract_pdf_text_with_ocr(file_path: Path) -> str:
    try:
        import fitz
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ValueError(
            "Dependências de OCR para PDF não estão instaladas. "
            "Instale PyMuPDF para renderizar páginas e use OCR em PDFs escaneados."
        ) from exc

    page_texts: list[str] = []
    try:
        with fitz.open(file_path) as document:
            for index, page in enumerate(document):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                text = pytesseract.image_to_string(image, lang="por+eng")
                if text.strip():
                    page_texts.append(f"[Página {index + 1}]\n{text.strip()}")
        return "\n\n".join(page_texts).strip()
    except Exception as exc:
        logger.error("Erro no OCR do PDF %s: %s", file_path, exc)
        raise


def extract_docx_text(file_path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ValueError("python-docx não instalado.") from exc

    try:
        document = docx.Document(file_path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n".join(paragraphs).strip()
    except Exception as exc:
        logger.error("Erro ao extrair DOCX %s: %s", file_path, exc)
        raise


def extract_xlsx_text(file_path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("openpyxl não instalado.") from exc

    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"--- Planilha: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]
                if any(values):
                    lines.append("\t".join(values))
        return "\n".join(lines).strip()
    except Exception as exc:
        logger.error("Erro ao extrair XLSX %s: %s", file_path, exc)
        raise


def extract_plain_text(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8", errors="replace").strip()
    except Exception as exc:
        logger.error("Erro ao ler arquivo texto %s: %s", file_path, exc)
        raise


def extract_image_text_with_ocr(file_path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise ValueError("Dependências de OCR de imagem não estão instaladas.") from exc

    try:
        image = Image.open(file_path)
        return pytesseract.image_to_string(image, lang="por+eng").strip()
    except Exception as exc:
        logger.error("Erro no OCR da imagem %s: %s", file_path, exc)
        raise


def write_extracted_text(output_path: Path, text: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        output_path.write_text(text or "", encoding="utf-8")
    except Exception as exc:
        logger.error("Falha ao salvar texto extraído em %s: %s", output_path, exc)
        raise

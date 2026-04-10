import logging
import io

from fastapi import APIRouter, HTTPException, UploadFile, File, BackgroundTasks
from starlette.concurrency import run_in_threadpool

from backend.models.schemas import FileGenerateRequest, FileGenerateResponse
from backend.models.schemas import (
    SessionFileListResponse,
    SessionFilesDeleteResponse,
    SessionFileUploadResponse,
)
from backend.services.file_service import generate_file
from backend.services.session_extraction_service import process_uploaded_file, recover_pending_session_files
from backend.services.session_gc_service import cleanup_expired_sessions
from backend.services.session_file_service import (
    SessionFileError,
    SessionLimitError,
    SessionNotFoundError,
    delete_session_dir,
    list_session_files,
    save_uploaded_file,
    session_exists,
    touch_session,
    validate_session_id,
)

router = APIRouter()
logger = logging.getLogger(__name__)


@router.post("/files", response_model=FileGenerateResponse)
async def files_endpoint(req: FileGenerateRequest):
    """Gera PDF, DOCX, XLSX ou PPTX e faz upload ao Supabase."""
    try:
        url, message = await generate_file(req.type, req.title, req.content)
        return FileGenerateResponse(url=url, message=message)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/session-files", response_model=SessionFileUploadResponse)
async def upload_session_file(
    background_tasks: BackgroundTasks,
    session_id: str,
    file: UploadFile = File(...),
):
    """Salva arquivo efêmero da sessão e dispara extração assíncrona."""
    try:
        cleanup_expired_sessions()
        validated_session_id = validate_session_id(session_id)
        content = await file.read()
        entry = save_uploaded_file(
            validated_session_id,
            file.filename or "arquivo",
            content,
            file.content_type or "application/octet-stream",
        )
        background_tasks.add_task(process_uploaded_file, validated_session_id, entry["file_id"])
        logger.info(
            "Upload efêmero recebido para sessão %s: %s",
            validated_session_id,
            entry["original_name"],
        )
        return SessionFileUploadResponse(
            session_id=validated_session_id,
            file_id=entry["file_id"],
            original_name=entry["original_name"],
            size_bytes=entry["size_bytes"],
            status="processing",
            message="Arquivo recebido e enviado para processamento em background.",
        )
    except SessionLimitError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except SessionFileError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.error("Falha no upload efêmero da sessão %s: %s", session_id, exc)
        raise HTTPException(status_code=500, detail="Erro interno ao salvar arquivo da sessão.") from exc


@router.get("/session-files/{session_id}", response_model=SessionFileListResponse)
async def get_session_files(session_id: str):
    """Lista os arquivos efêmeros da sessão."""
    try:
        cleanup_expired_sessions()
        validated_session_id = validate_session_id(session_id)
        if not session_exists(validated_session_id):
            return SessionFileListResponse(session_id=validated_session_id, files=[])
        await run_in_threadpool(recover_pending_session_files, validated_session_id)
        touch_session(validated_session_id)
        return SessionFileListResponse(
            session_id=validated_session_id,
            files=list_session_files(validated_session_id),
        )
    except SessionNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except SessionFileError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.error("Falha ao listar arquivos da sessão %s: %s", session_id, exc)
        raise HTTPException(status_code=500, detail="Erro interno ao listar arquivos da sessão.") from exc


@router.delete("/session-files/{session_id}", response_model=SessionFilesDeleteResponse)
async def delete_session_files(session_id: str):
    """Remove todos os arquivos e o manifesto da sessão efêmera."""
    try:
        cleanup_expired_sessions()
        validated_session_id = validate_session_id(session_id)
        delete_session_dir(validated_session_id)
        return SessionFilesDeleteResponse(
            session_id=validated_session_id,
            deleted=True,
            message="Sessão efêmera removida com sucesso.",
        )
    except SessionNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except SessionFileError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.error("Falha ao remover sessão %s: %s", session_id, exc)
        raise HTTPException(status_code=500, detail="Erro interno ao remover sessão.") from exc

@router.post("/extract-text")
async def extract_text_endpoint(file: UploadFile = File(...)):
    """Extrai texto de PDFs, DOCX ou XLSX para uso no chat."""
    filename = file.filename.lower()
    
    try:
        content = await file.read()
        
        if filename.endswith(".pdf"):
            import PyPDF2
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() for page in pdf_reader.pages if page.extract_text())
            return {"text": text}
            
        elif filename.endswith(".docx"):
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            return {"text": text}
            
        elif filename.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            text_lines = []
            for sheet in wb.worksheets:
                text_lines.append(f"--- Planilha: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    # Filter out purely None rows
                    row_vals = [str(cell) if cell is not None else "" for cell in row]
                    if any(row_vals):
                        text_lines.append("\t".join(row_vals))
            return {"text": "\n".join(text_lines)}
            
        else:
            # Fallback for txt, csv, md
            return {"text": content.decode("utf-8", errors="replace")}
            
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao extrair texto: {str(e)}")
